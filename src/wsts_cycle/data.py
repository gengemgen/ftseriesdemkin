"""Extraction and loading utilities for the WSTS workbook."""

from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

from wsts_cycle.paths import PROCESSED_DATA_DIR, RAW_DATA_DIR

RAW_XLSX = RAW_DATA_DIR / "WSTS-Historical-Billings-Report-May_2026.xlsx"

REGIONS = ("Americas", "Europe", "Japan", "Asia Pacific", "Worldwide")
SHEETS = {"Monthly Data": "monthly", "3MMA": "3mma"}
UNITS_TEXT = "All numbers are in 1000 US$"


def _parse_sheet(ws: Any, series_name: str) -> pd.DataFrame:
    """Parse one worksheet into tidy monthly observations.

    Parameters
    ----------
    ws : openpyxl.worksheet.worksheet.Worksheet
        Workbook sheet containing one WSTS series.
    series_name : {"monthly", "3mma"}
        Name assigned to observations from the worksheet.

    Returns
    -------
    pandas.DataFrame
        Tidy observations with date, region, series, and billings columns.
    """

    rows = list(ws.iter_rows(values_only=True))
    units = str(rows[2][0]).rstrip(".")
    if units != UNITS_TEXT:
        raise ValueError(f"units text changed: {rows[2][0]!r}")

    header = rows[3]
    months = [
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    ]
    if list(header[1:13]) != months:
        raise ValueError(f"month header changed: {header[1:13]}")

    records = []
    row_number = 4
    while row_number < len(rows):
        first = rows[row_number][0]
        if first is None:
            row_number += 1
            continue

        year = int(first)
        block = rows[row_number + 1 : row_number + 1 + len(REGIONS)]
        observed_regions = [row[0] for row in block]
        if observed_regions != list(REGIONS):
            raise ValueError(
                f"region block for {year} changed: {observed_regions}"
            )

        for region_row in block:
            region = region_row[0]
            for month, value in enumerate(region_row[1:13], start=1):
                if value is None:
                    continue
                date = pd.Timestamp(year=year, month=month, day=1)
                if series_name == "3mma" and value == 0:
                    records.append((date, region, series_name, None))
                    continue
                records.append((date, region, series_name, float(value)))
        row_number += 1 + len(REGIONS)

    frame = pd.DataFrame.from_records(
        records,
        columns=["date", "region", "series", "billings_kusd"],
    )

    placeholders = frame[frame["billings_kusd"].isna()]
    if not placeholders.empty:
        last_observation = frame.dropna(subset=["billings_kusd"])["date"].max()
        internal_zeros = placeholders[placeholders["date"] <= last_observation]
        if not internal_zeros.empty:
            raise ValueError(
                "zero values inside the published range:\n"
                f"{internal_zeros}"
            )
        frame = frame.dropna(subset=["billings_kusd"])
    return frame


def extract(raw_xlsx: Path = RAW_XLSX) -> dict[str, pd.DataFrame]:
    """Extract the monthly and three-month-average WSTS series.

    Parameters
    ----------
    raw_xlsx : pathlib.Path, optional
        Source WSTS workbook.

    Returns
    -------
    dict[str, pandas.DataFrame]
        Tidy observations and wide monthly and 3MMA tables.
    """

    workbook = openpyxl.load_workbook(raw_xlsx, read_only=True, data_only=True)
    try:
        parts = [
            _parse_sheet(workbook[sheet], series)
            for sheet, series in SHEETS.items()
        ]
    finally:
        workbook.close()

    tidy = pd.concat(parts, ignore_index=True).sort_values(
        ["series", "region", "date"], ignore_index=True
    )
    wide = {
        series: (
            tidy[tidy["series"] == series]
            .pivot(index="date", columns="region", values="billings_kusd")
            .loc[:, REGIONS]
        )
        for series in ("monthly", "3mma")
    }
    return {"tidy": tidy, "monthly": wide["monthly"], "3mma": wide["3mma"]}


def write_csvs(
    frames: dict[str, pd.DataFrame], outdir: Path = PROCESSED_DATA_DIR
) -> None:
    """Write extracted WSTS tables to CSV files.

    Parameters
    ----------
    frames : dict[str, pandas.DataFrame]
        Output returned by :func:`extract`.
    outdir : pathlib.Path, optional
        Directory for processed CSV files.
    """

    outdir.mkdir(parents=True, exist_ok=True)
    frames["tidy"].to_csv(outdir / "wsts_tidy_long.csv", index=False)
    frames["monthly"].to_csv(outdir / "wsts_wide_monthly.csv")
    frames["3mma"].to_csv(outdir / "wsts_wide_3mma.csv")


def _load_wide(csv_name: str) -> pd.DataFrame:
    """Load one processed wide table with a monthly date index."""

    frame = pd.read_csv(
        PROCESSED_DATA_DIR / csv_name,
        index_col="date",
        parse_dates=True,
    )
    frame.index.freq = None
    return frame


def load_monthly(region: str = "Worldwide") -> pd.Series:
    """Load a gap-free monthly billings series.

    Parameters
    ----------
    region : {"Americas", "Europe", "Japan", "Asia Pacific", "Worldwide"}
        Geographic series to load.

    Returns
    -------
    pandas.Series
        Monthly billings in thousands of US dollars.

    Raises
    ------
    ValueError
        If the region is unknown or the monthly series contains gaps.
    """

    if region not in REGIONS:
        raise ValueError(f"unknown region {region!r}")
    endog = _load_wide("wsts_wide_monthly.csv")[region].dropna().asfreq("MS")
    if endog.isna().any():
        raise ValueError(f"gaps in monthly series for {region}")

    suffix = region.lower().replace(" ", "_")
    endog.name = f"billings_kusd_{suffix}"
    return endog


def load_3mma(region: str = "Worldwide") -> pd.Series:
    """Load a gap-free WSTS three-month moving-average series.

    Parameters
    ----------
    region : {"Americas", "Europe", "Japan", "Asia Pacific", "Worldwide"}
        Geographic series to load.

    Returns
    -------
    pandas.Series
        Published 3MMA billings in thousands of US dollars.
    """

    if region not in REGIONS:
        raise ValueError(f"unknown region {region!r}")
    endog = _load_wide("wsts_wide_3mma.csv")[region].dropna().asfreq("MS")
    if endog.isna().any():
        raise ValueError(f"gaps in 3MMA series for {region}")
    suffix = region.lower().replace(" ", "_")
    endog.name = f"billings_3mma_kusd_{suffix}"
    return endog
